import io
from docx import Document as Docx
from ingestion import Document, Chunk, extract
from narrative_intelligence import (
    ONTOLOGY, NarrativeFact, SemanticFactCandidate, classify_source_role,
    extract_narrative_facts, validate_semantic_fact,
)
from narrative_claims import GroundedNarrativeClaimExtractor


def doc(name,text,kind="artifact",locator="paragraph 1"):
    return Document(name,kind,(Chunk(name,locator,text),))

def facts(text,name="evidence.txt",kind="artifact"):
    return extract_narrative_facts([doc(name,text,kind)])

def test_ontology_covers_required_security_domains():
    required={"identity_access","mfa","joiner_mover_leaver","access_review","privileged_access","authentication","encryption_at_rest","encryption_in_transit","key_management","network_security","firewalls","segmentation","endpoint_security","logging","monitoring","alerting","incident_response","vulnerability_management","patching","change_management","backup","business_continuity","disaster_recovery","vendor_management","risk_management","asset_management","data_inventory","data_deletion","physical_security","security_awareness","background_screening"}
    assert required <= set(ONTOLOGY)

def test_policy_requirement_is_grounded_but_not_operating_evidence():
    f=facts("MFA must be required for production access.","security.txt")
    x=next(v for v in f if v.concept=="mfa" and v.attribute=="requirement")
    assert x.source_role in {"POLICY","UNKNOWN"}
    assert x.provenance.excerpt=="MFA must be required for production access."
    assert x.claim_type=="REQUIREMENT"

def test_content_first_role_beats_misleading_filename():
    d=doc("random-export.csv","Information Security Policy. All workforce access must use MFA.",kind="artifact",locator="row 1")
    assert classify_source_role(d)=="POLICY"

def test_negative_mfa_is_preserved_as_negative_not_affirmative_claim():
    f=facts("MFA is not enforced for legacy VPN access.")
    x=next(v for v in f if v.concept=="mfa" and v.attribute=="enforcement_evidence")
    assert "NEGATED" in x.qualifiers and x.normalized_value is False
    claims=GroundedNarrativeClaimExtractor().extract([doc("x.txt","MFA is not enforced for legacy VPN access.")])
    assert not any(c.topic=="mfa" and c.attributes.get("enforcement_evidence") is True for c in claims)

def test_planned_control_does_not_become_current_operation():
    f=facts("MFA will be enforced for production systems next quarter.")
    x=next(v for v in f if v.concept=="mfa")
    assert "PLANNED_OR_FUTURE" in x.qualifiers
    assert x.normalized_value=="PLANNED"

def test_exception_is_preserved():
    f=facts("MFA is enforced for production except break-glass accounts.")
    assert any("EXCEPTION" in x.qualifiers for x in f)

def test_conditional_language_is_preserved():
    f=facts("Endpoint protection should be deployed where feasible.")
    assert any("CONDITIONAL" in x.qualifiers for x in f)

def test_multiple_sentences_get_exact_supporting_excerpt():
    d=doc("a.txt","Noise sentence. Administrative access is restricted to authorized personnel. More noise.",kind="policy")
    f=extract_narrative_facts([d])
    x=next(v for v in f if v.attribute=="privilege_restricted")
    assert x.provenance.excerpt=="Administrative access is restricted to authorized personnel."

def test_docx_heading_and_paragraph_segmentation_is_preserved():
    d=Docx(); d.add_heading("Access Control",level=1); d.add_paragraph("MFA must be required for production access.")
    b=io.BytesIO(); d.save(b)
    parsed=extract("policy.docx",b.getvalue())
    para=next(c for c in parsed.chunks if "MFA" in c.text)
    assert para.heading=="Access Control" and para.segment_type=="paragraph"

def test_semantic_candidate_must_quote_exact_locator():
    d=doc("x.txt","MFA is enforced for production.")
    good=SemanticFactCandidate("mfa","enforcement_evidence",True,"OBSERVATION","MFA is enforced for production.","paragraph 1",.91,"CONFIGURATION")
    assert validate_semantic_fact(good,d,{"mfa"}).accepted
    bad=SemanticFactCandidate("mfa","enforcement_evidence",True,"OBSERVATION","MFA is enforced everywhere.","paragraph 1",.91,"CONFIGURATION")
    assert not validate_semantic_fact(bad,d,{"mfa"}).accepted

def test_semantic_candidate_rejects_wrong_locator_and_concept():
    d=doc("x.txt","MFA is enforced for production.")
    wrong=SemanticFactCandidate("mfa","enforcement_evidence",True,"OBSERVATION","MFA is enforced for production.","page 99",.91,"CONFIGURATION")
    assert not validate_semantic_fact(wrong,d,{"mfa"}).accepted
    concept=SemanticFactCandidate("risk_management","x",True,"OBSERVATION","MFA is enforced for production.","paragraph 1",.91,"CONFIGURATION")
    assert not validate_semantic_fact(concept,d,{"mfa"}).accepted

def test_semantic_candidate_rejects_verdict_language():
    d=doc("x.txt","The control passed according to the reviewer.")
    c=SemanticFactCandidate("mfa","x","control passed","OBSERVATION","The control passed according to the reviewer.","paragraph 1",.9,"REPORT")
    assert not validate_semantic_fact(c,d,{"mfa"}).accepted

def test_major_domain_examples_extract_grounded_facts():
    samples=[
      ("Logging Policy. Security logs must be centrally retained.","logging_requirement"),
      ("Endpoint Security Policy. Endpoint EDR protection is required on production assets.","endpoint_requirement"),
      ("Incident severity criteria define when an incident is declared.","incident_criteria"),
      ("Vulnerability scanning scope must include all production assets.","scan_scope_requirement"),
      ("Change approval is required for production configuration changes.","configuration_baseline_requirement"),
      ("Media disposal must render protected data unrecoverable.","media_disposal_requirement"),
    ]
    for text,attr in samples:
        assert any(f.attribute==attr and f.provenance.excerpt for f in facts(text,"policy.txt","policy")), (text,attr)

def test_tabular_header_is_not_promoted_to_operating_fact():
    d=Document("edr.csv","artifact",(Chunk("edr.csv","row 1","Asset | EDR Installed | Status"),))
    assert extract_narrative_facts([d]) == []

def test_semantic_provider_candidates_are_validated_not_trusted():
    from narrative_intelligence import extract_validated_semantic_facts
    d=doc("x.txt","MFA is enforced for production.")
    class Provider:
        def extract_candidates(self, document, allowed_concepts):
            return [
                SemanticFactCandidate("mfa","enforcement_evidence",True,"OBSERVATION","MFA is enforced for production.","paragraph 1",.91,"CONFIGURATION"),
                SemanticFactCandidate("mfa","enforcement_evidence",True,"OBSERVATION","invented quote","paragraph 1",.99,"CONFIGURATION"),
            ]
    accepted,rejected=extract_validated_semantic_facts(d,Provider(),{"mfa"})
    assert len(accepted)==1 and len(rejected)==1
