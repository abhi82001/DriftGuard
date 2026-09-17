#!/usr/bin/env python3
"""Realistic USER_ACCESS_REVIEW fixtures for CP007 Phase 1 tests.

Two artifacts, built in memory so no binary is committed:

`campaign_workbook()` - a full quarterly campaign export with a summary sheet, a
per-entitlement detail sheet and an exclusions sheet. It covers population
reconciliation, decision reconciliation, completed and open remediation,
exclusions, privileged and service-account populations, and missing supporting
evidence (there is no post-change confirmation column). It also contains one
deliberate contradiction: the summary states 8 excluded accounts while the
exclusions sheet lists 9.

`summary_csv()` - a small clean CSV summary with no contradiction, used to prove
the arithmetic reconciles when the artifact is internally consistent. It also
reproduces the worked example: 9 revocations, 8 completed, 1 left open.

Alongside them, `real_workbook()` reads a committed customer-shaped export and
`edge_workbook()` builds detail-only sheets for the row-level cases the tidy
fixtures never produced: a decision value the extractor cannot place, an
identity value that matches no known type, and a post-change confirmation
column that is present without confirming everything.
"""

from __future__ import annotations

import io
from pathlib import Path

FILES_DIR = Path(__file__).resolve().parent / "files"

WORKBOOK_FILENAME = "q3_2026_user_access_review.xlsx"
CSV_FILENAME = "q2_2026_access_review_summary.csv"

# Summary sheet, as stated by the campaign owner.
SUMMARY_SOURCE_POPULATION = 120
SUMMARY_EXCLUDED = 8                  # contradicts the exclusions sheet (9 rows)
SUMMARY_REVIEWED = 112
SUMMARY_RETAIN = 90
SUMMARY_MODIFY = 8
SUMMARY_REVOKE = 9
SUMMARY_INVESTIGATE = 5
SUMMARY_REMEDIATION_COMPLETED = 8
SUMMARY_REMEDIATION_OPEN = 9
SUMMARY_EXCEPTIONS = 2
EXCLUSION_ROWS = 9                    # the deliberate contradiction

SUMMARY_SHEET = "Campaign Summary"
DETAIL_SHEET = "Review Detail"
EXCLUSIONS_SHEET = "Exclusions"

EXCLUDED_CELL = "B11"                 # where the contradicted value sits
EXCLUSIONS_RANGE = "rows 2-10"
DETAIL_RANGE = "rows 2-113"

_SUMMARY_ROWS: list[list] = [
    ["Q3 2026 User Access Review - Campaign Summary"],
    [],
    ["Campaign Name", "Q3-2026-UAR-PROD"],
    ["Campaign Status", "Closed"],
    ["Review Period Start", "2026-07-01"],
    ["Review Period End", "2026-09-30"],
    ["Review Completed Date", "2026-10-05"],
    ["Reviewer", "D. Okonkwo (System Owner)"],
    [],
    ["Source Population (accounts extracted)", SUMMARY_SOURCE_POPULATION],
    ["Excluded Accounts", SUMMARY_EXCLUDED],
    ["Reviewed / Certified Accounts", SUMMARY_REVIEWED],
    [],
    ["Decision: Retain", SUMMARY_RETAIN],
    ["Decision: Modify", SUMMARY_MODIFY],
    ["Decision: Revoke", SUMMARY_REVOKE],
    ["Decision: Investigate", SUMMARY_INVESTIGATE],
    [],
    ["Remediation Completed", SUMMARY_REMEDIATION_COMPLETED],
    ["Remediation Open", SUMMARY_REMEDIATION_OPEN],
    ["Open Exceptions", SUMMARY_EXCEPTIONS],
]

_DETAIL_HEADERS = ["Account ID", "Display Name", "Identity Type", "System",
                   "Reviewer", "Review Decision", "Remediation Ticket",
                   "Remediation Status"]

# Identity mix across the 112 reviewed accounts.
_IDENTITY_MIX = (("Employee", 80), ("Contractor", 15), ("Service Account", 10),
                 ("Privileged Admin", 7))
_DECISION_MIX = (("Retain", SUMMARY_RETAIN), ("Modify", SUMMARY_MODIFY),
                 ("Revoke", SUMMARY_REVOKE), ("Investigate", SUMMARY_INVESTIGATE))


def _expand(mix) -> list[str]:
    out: list[str] = []
    for value, count in mix:
        out.extend([value] * count)
    return out


def _detail_rows() -> list[list]:
    identities = _expand(_IDENTITY_MIX)
    decisions = _expand(_DECISION_MIX)
    assert len(identities) == len(decisions) == SUMMARY_REVIEWED

    rows: list[list] = [list(_DETAIL_HEADERS)]
    completed_left = SUMMARY_REMEDIATION_COMPLETED
    ticket = 4100
    for index, (identity, decision) in enumerate(zip(identities, decisions), start=1):
        reference = status = ""
        if decision in ("Modify", "Revoke"):
            ticket += 1
            reference = f"CHG-{ticket}"
            if completed_left:
                status, completed_left = "Completed", completed_left - 1
            else:
                status = "Open"
        rows.append([
            f"ACC-{index:04d}", f"user{index:03d}@acme.test", identity,
            "prod-erp", "D. Okonkwo", decision, reference, status,
        ])
    return rows


def _exclusion_rows() -> list[list]:
    reasons = [
        "Break-glass account, reviewed under separate privileged process",
        "Disabled before the review period opened",
        "Vendor support account governed by contract review",
        "Test account in non-production tenant",
        "Account pending deletion at period start",
        "Integration account owned by the platform team",
        "Duplicate identity merged during the period",
        "Read-only monitoring account",
        "Account created after the population extraction date",
    ]
    assert len(reasons) == EXCLUSION_ROWS
    rows: list[list] = [["Account ID", "Identity Type", "Exclusion Reason"]]
    for index, reason in enumerate(reasons, start=1):
        rows.append([f"EXC-{index:04d}", "Service Account", reason])
    return rows


def campaign_workbook() -> bytes:
    """Build the full XLSX campaign export (contains the deliberate contradiction)."""
    return _build_workbook((
        (SUMMARY_SHEET, _SUMMARY_ROWS),
        (DETAIL_SHEET, _detail_rows()),
        (EXCLUSIONS_SHEET, _exclusion_rows()),
    ))


def _build_workbook(sheets) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value not in (None, ""):
                    ws.cell(row=r, column=c, value=value)
    wb.remove(wb.worksheets[0])          # drop the default empty sheet
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


CSV_SOURCE_POPULATION = 80
CSV_EXCLUDED = 5
CSV_REVIEWED = 75
CSV_RETAIN = 60
CSV_MODIFY = 0
CSV_REVOKE = 9
CSV_INVESTIGATE = 6
CSV_REMEDIATION_COMPLETED = 8
CSV_OPEN_REMEDIATION = 1                 # 9 flagged for change - 8 completed


def summary_csv() -> bytes:
    """A clean, internally consistent CSV summary (no contradiction)."""
    lines = [
        "Metric,Value",
        "Campaign Name,Q2-2026-UAR-FIN",
        "Campaign Status,Closed",
        "Review Period Start,2026-04-01",
        "Review Period End,2026-06-30",
        f"Source Population,{CSV_SOURCE_POPULATION}",
        f"Excluded Accounts,{CSV_EXCLUDED}",
        f"Reviewed Accounts,{CSV_REVIEWED}",
        f"Decision: Retain,{CSV_RETAIN}",
        f"Decision: Modify,{CSV_MODIFY}",
        f"Decision: Revoke,{CSV_REVOKE}",
        f"Decision: Investigate,{CSV_INVESTIGATE}",
        f"Remediation Completed,{CSV_REMEDIATION_COMPLETED}",
    ]
    return ("\n".join(lines) + "\n").encode("utf-8")


# --------------------------------------------------------------------------
# A fully reconciled campaign: every check reconciles and every supporting
# element the validator looks for is present, so the report shows 0 exceptions.
CLEAN_FILENAME = "q1_2026_user_access_review_clean.xlsx"
CLEAN_SOURCE_POPULATION = 50
CLEAN_EXCLUDED = 4
CLEAN_REVIEWED = 46
CLEAN_RETAIN = 40
CLEAN_MODIFY = 3
CLEAN_REVOKE = 2
CLEAN_INVESTIGATE = 1
CLEAN_REMEDIATION_COMPLETED = CLEAN_MODIFY + CLEAN_REVOKE     # nothing left open

_CLEAN_SUMMARY_ROWS: list[list] = [
    ["Q1 2026 User Access Review - Campaign Summary"],
    [],
    ["Campaign Name", "Q1-2026-UAR-CORE"],
    ["Campaign Status", "Closed"],
    ["Review Period Start", "2026-01-01"],
    ["Review Period End", "2026-03-31"],
    [],
    ["Source Population (accounts extracted)", CLEAN_SOURCE_POPULATION],
    ["Excluded Accounts", CLEAN_EXCLUDED],
    ["Reviewed / Certified Accounts", CLEAN_REVIEWED],
    [],
    ["Decision: Retain", CLEAN_RETAIN],
    ["Decision: Modify", CLEAN_MODIFY],
    ["Decision: Revoke", CLEAN_REVOKE],
    ["Decision: Investigate", CLEAN_INVESTIGATE],
    [],
    ["Remediation Completed", CLEAN_REMEDIATION_COMPLETED],
    ["Remediation Open", 0],
    ["Open Exceptions", 0],
]

_CLEAN_IDENTITY_MIX = (("Employee", 30), ("Contractor", 8),
                       ("Service Account", 5), ("Privileged Admin", 3))
_CLEAN_DECISION_MIX = (("Retain", CLEAN_RETAIN), ("Modify", CLEAN_MODIFY),
                       ("Revoke", CLEAN_REVOKE), ("Investigate", CLEAN_INVESTIGATE))


def _clean_detail_rows() -> list[list]:
    identities = _expand(_CLEAN_IDENTITY_MIX)
    decisions = _expand(_CLEAN_DECISION_MIX)
    assert len(identities) == len(decisions) == CLEAN_REVIEWED
    rows: list[list] = [["Account ID", "Display Name", "Identity Type", "System",
                         "Reviewer", "Review Decision", "Remediation Ticket",
                         "Remediation Status", "Post-Change Confirmed"]]
    ticket = 2200
    for index, (identity, decision) in enumerate(zip(identities, decisions), start=1):
        reference = status = confirmed = ""
        if decision in ("Modify", "Revoke"):
            ticket += 1
            reference, status, confirmed = f"CHG-{ticket}", "Completed", "Yes"
        rows.append([f"ACC-{index:04d}", f"user{index:03d}@acme.test", identity,
                     "prod-crm", "R. Mehta", decision, reference, status, confirmed])
    return rows


def _clean_exclusion_rows() -> list[list]:
    reasons = ["Break-glass account under separate privileged review",
               "Disabled before the review period opened",
               "Test account in non-production tenant",
               "Duplicate identity merged during the period"]
    assert len(reasons) == CLEAN_EXCLUDED
    rows: list[list] = [["Account ID", "Identity Type", "Exclusion Reason"]]
    for index, reason in enumerate(reasons, start=1):
        rows.append([f"EXC-{index:04d}", "Service Account", reason])
    return rows


def clean_workbook() -> bytes:
    """A campaign export in which every validation check reconciles."""
    return _build_workbook((
        (SUMMARY_SHEET, _CLEAN_SUMMARY_ROWS),
        (DETAIL_SHEET, _clean_detail_rows()),
        (EXCLUSIONS_SHEET, _clean_exclusion_rows()),
    ))


# --------------------------------------------------------------------------
# A real customer-shaped export, committed as a file rather than built in
# memory: one detail sheet, no summary, and cell values that the in-memory
# fixtures were all too tidy to contain.
REAL_FILENAME = "Q3_Access_Review.xlsx"


def real_workbook() -> bytes:
    return (FILES_DIR / REAL_FILENAME).read_bytes()


# --------------------------------------------------------------------------
# Single-sheet detail exports for the row-level edge cases: values the
# extractor cannot place, and a post-change confirmation column that is
# present without confirming everything (or anything).
EDGE_FILENAME = "q4_2026_access_review_detail.xlsx"
EDGE_SHEET = "Access Review"
EDGE_HEADERS = ["Account ID", "Identity Type", "Reviewer", "Review Decision",
                "Remediation Ticket", "Remediation Status",
                "Post-Change Confirmed"]


def edge_workbook(rows: list[list]) -> bytes:
    """A detail-only access review export built from the given data rows."""
    return _build_workbook(((EDGE_SHEET, [list(EDGE_HEADERS)] + [list(r) for r in rows]),))


# Two rows the extractor cannot place: an "Excluded" disposition that is not a
# reviewer decision, and a row left blank. Neither may be dropped.
UNRECOGNIZED_DECISION_ROWS = [
    ["ACC-0001", "Employee", "L. Martin", "Retain", "", "", ""],
    ["ACC-0002", "Privileged", "R. Patel", "Excluded", "", "", ""],
    ["ACC-0003", "Service Account", "Security", "", "", "", ""],
]

# "Robot" matches no known identity type and must not become one; the blank in
# the third row is an absent value rather than an unrecognised one.
UNMAPPED_IDENTITY_ROWS = [
    ["ACC-0001", "Employee", "L. Martin", "Retain", "", "", ""],
    ["ACC-0002", "Robot", "R. Patel", "Retain", "", "", ""],
    ["ACC-0003", "", "Security", "Retain", "", "", ""],
]

# The confirmation column exists and is entirely empty, while the rows plainly
# record remediation.
EMPTY_CONFIRMATION_ROWS = [
    ["ACC-0001", "Employee", "L. Martin", "Retain", "", "", ""],
    ["ACC-0002", "Privileged", "R. Patel", "Revoke", "UAR-9001", "Closed", ""],
    ["ACC-0003", "Contractor", "Vendor Owner", "Revoke", "UAR-9002", "Open", ""],
]

# One row affirms confirmation; the other carries a timestamp, which is not an
# affirmative token and is not guessed at.
PARTIAL_CONFIRMATION_ROWS = [
    ["ACC-0001", "Employee", "L. Martin", "Retain", "", "", ""],
    ["ACC-0002", "Privileged", "R. Patel", "Revoke", "UAR-9001", "Closed", "Yes"],
    ["ACC-0003", "Contractor", "Vendor Owner", "Revoke", "UAR-9002", "Closed",
     "2026-10-07 08:52 UTC"],
]

# Nothing was remediated at all, so the confirmation column has nothing to cover.
NO_REMEDIATION_ROWS = [
    ["ACC-0001", "Employee", "L. Martin", "Retain", "", "", ""],
    ["ACC-0002", "Privileged", "R. Patel", "Retain", "", "", ""],
]


# A tabular upload that is NOT an access review: it must stay UNCLASSIFIED.
NON_REVIEW_FILENAME = "idp_mfa_configuration_export.csv"
NON_REVIEW_CSV = (
    "application,authentication policy,state\n"
    "identity provider,multi-factor authentication,enforced\n"
    "production console,multi-factor authentication,enabled\n"
).encode()
