#!/usr/bin/env python3
"""Structure-preserving XLSX/CSV reading for evidence analysis (CP007 Phase 1).

`ingestion.py` flattens a spreadsheet into prose chunks so the claim extractor
can read it as text. Evidence analysis needs the opposite: sheets, header rows
and addressable cells, so every extracted fact can point back at the cell it
came from. This module adds that view without changing ingestion.

Uploaded content is data. Nothing here is executed, evaluated, or interpreted
as an instruction.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from typing import Any, Optional

from ingestion import MAX_BYTES, IngestionError

TABULAR_EXTENSIONS = (".xlsx", ".csv")
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 5000
MAX_COLUMNS = 64
MIN_HEADER_CELLS = 3


class TabularError(IngestionError):
    """A tabular upload could not be read as a workbook."""


def column_letter(index: int) -> str:
    """1 -> A, 2 -> B, 27 -> AA (spreadsheet column addressing)."""
    if index < 1:
        raise ValueError("column index is 1-based")
    out = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        out = chr(ord("A") + remainder) + out
    return out


@dataclass(frozen=True)
class Cell:
    row: int              # 1-based
    column: int           # 1-based
    value: Any

    @property
    def ref(self) -> str:
        return f"{column_letter(self.column)}{self.row}"

    @property
    def text(self) -> str:
        if self.value is None:
            return ""
        if isinstance(self.value, str):
            return " ".join(self.value.split())
        return str(self.value)

    @property
    def empty(self) -> bool:
        return not self.text.strip()


@dataclass(frozen=True)
class Row:
    number: int           # 1-based
    cells: tuple[Cell, ...]

    @property
    def filled(self) -> tuple[Cell, ...]:
        return tuple(c for c in self.cells if not c.empty)

    def cell(self, column: int) -> Optional[Cell]:
        return next((c for c in self.cells if c.column == column), None)


@dataclass(frozen=True)
class Sheet:
    filename: str
    name: str
    rows: tuple[Row, ...]
    header_row: Optional[int] = None
    headers: tuple[str, ...] = ()          # normalized header text, in column order
    header_columns: tuple[int, ...] = ()   # 1-based column index per header

    @property
    def data_rows(self) -> tuple[Row, ...]:
        """Non-empty rows after the header row (or all rows when there is none)."""
        start = (self.header_row or 0) + 1
        return tuple(r for r in self.rows if r.number >= start and r.filled)

    def column_of(self, *header_names: str) -> Optional[int]:
        """1-based index of the first column whose header matches exactly."""
        wanted = {normalize(h) for h in header_names}
        for text, index in zip(self.headers, self.header_columns):
            if text in wanted:
                return index
        return None

    def header_matching(self, *fragments: str) -> Optional[int]:
        """1-based index of the first column whose header contains a fragment."""
        for text, index in zip(self.headers, self.header_columns):
            if any(f in text for f in fragments):
                return index
        return None


@dataclass(frozen=True)
class Workbook:
    filename: str
    sheets: tuple[Sheet, ...]


def normalize(text: Any) -> str:
    """Lowercase, whitespace-collapsed, trailing-punctuation-free label text."""
    if text is None:
        return ""
    if not isinstance(text, str):
        text = str(text)
    return " ".join(text.split()).strip().strip(":").strip().lower()


def is_tabular(filename: str) -> bool:
    return filename.lower().endswith(TABULAR_EXTENSIONS)


def read_tabular(filename: str, data: bytes) -> Workbook:
    """Read an XLSX or CSV upload into addressable sheets."""
    lower = filename.lower()
    if not is_tabular(filename):
        raise TabularError(f"{filename} is not a tabular file")
    if not data:
        raise TabularError(f"{filename} is empty")
    if len(data) > MAX_BYTES:
        raise TabularError(f"{filename} exceeds {MAX_BYTES // (1024 * 1024)} MB")

    try:
        sheets = _xlsx(filename, data) if lower.endswith(".xlsx") else _csv(filename, data)
    except TabularError:
        raise
    except Exception as exc:  # noqa: BLE001 - a malformed upload must not crash
        raise TabularError(f"could not read {filename}: {type(exc).__name__}") from exc

    sheets = tuple(s for s in sheets if s.rows)
    if not sheets:
        raise TabularError(f"no tabular content in {filename}")
    return Workbook(filename, sheets)


# ------------------------------------------------------------------- readers
def _rows_from_values(values: list[list[Any]]) -> tuple[Row, ...]:
    rows = []
    for number, raw in enumerate(values[:MAX_ROWS_PER_SHEET], start=1):
        cells = tuple(
            Cell(number, column, value)
            for column, value in enumerate(raw[:MAX_COLUMNS], start=1)
        )
        rows.append(Row(number, cells))
    while rows and not rows[-1].filled:
        rows.pop()
    return tuple(rows)


def _sheet(filename: str, name: str, values: list[list[Any]]) -> Sheet:
    rows = _rows_from_values(values)
    header_row, headers, columns = _detect_header(rows)
    return Sheet(filename, name, rows, header_row, headers, columns)


def _detect_header(rows: tuple[Row, ...]) -> tuple[Optional[int], tuple[str, ...], tuple[int, ...]]:
    """First row that looks like a header: several distinct text labels with data under it.

    A two-column label/value summary sheet deliberately does not qualify; it is
    read as key/value pairs instead.
    """
    for position, row in enumerate(rows):
        filled = row.filled
        if len(filled) < MIN_HEADER_CELLS:
            continue
        if not all(isinstance(c.value, str) for c in filled):
            continue
        texts = [normalize(c.text) for c in filled]
        if len(set(texts)) != len(texts):
            continue
        if not any(r.filled for r in rows[position + 1:]):
            continue
        return row.number, tuple(texts), tuple(c.column for c in filled)
    return None, (), ()


def _csv(filename: str, data: bytes) -> tuple[Sheet, ...]:
    text = data.decode("utf-8-sig", errors="replace")
    values = [list(row) for row in csv.reader(io.StringIO(text))]
    stem = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    return (_sheet(filename, stem, values),)


def _xlsx(filename: str, data: bytes) -> tuple[Sheet, ...]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in wb.worksheets[:MAX_SHEETS]:
            values = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheets.append(_sheet(filename, worksheet.title, values))
    finally:
        wb.close()
    return tuple(sheets)
